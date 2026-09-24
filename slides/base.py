"""
Primitivas do design system LAB365 × SENAI sobre python-pptx.

Tudo que desenha num slide passa por aqui: template oficial, tokens,
medição de texto com a fonte Lexend (para o texto nunca estourar a caixa),
formas, ícones Lucide recoloridos, cabeçalho e barra de destaque.

Unidades: todas as posições e tamanhos são em polegadas (float);
tamanhos de fonte em pontos.
"""

from __future__ import annotations

import difflib
import functools
import io
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence, Union

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# ── Caminhos e tokens ─────────────────────────────────────────────────────────

DIR = Path(__file__).resolve().parent
RAIZ_PROJETO = DIR.parent
TEMPLATE = DIR / "template" / "lab365.pptx"
DIR_ICONES = DIR / "icones"
DIR_FONTES = DIR / "fontes"
CAMINHO_TOKENS = DIR / "tokens.json"

EMU_POR_POL = 914400


def carregar_tokens(caminho: Union[str, Path] = CAMINHO_TOKENS) -> dict:
    """Lê o tokens.json do design system."""
    with open(caminho, encoding="utf-8") as f:
        return json.load(f)


TOKENS: dict = carregar_tokens()
CORES: dict = TOKENS["cores"]
GRID: dict = TOKENS["grid"]
FONTES: dict = TOKENS["fonte"]
ESTILOS: dict = TOKENS["estilos"]

LARGURA = GRID["largura"]
ALTURA = GRID["altura"]
MX = GRID["margem_x"]                 # margem lateral
LU = GRID["largura_util"]             # largura útil (MX .. MX+LU)
TOPO = GRID["topo_conteudo"]          # onde o conteúdo começa
BASE = GRID["base_conteudo"]          # onde o conteúdo termina
TOPO_NOTA = GRID["topo_nota"]         # barra de destaque inferior
ALT_NOTA = GRID["altura_nota"]
ESPACO = GRID["espaco"]               # respiro padrão entre blocos
RAIO = GRID["raio_card"]              # raio de card
RAIO_LINHA = GRID["raio_linha"]       # raio de linha de tabela

# Índices dos layouts do template oficial
LAYOUTS = {
    "capa": 0,          # ondas + faixa branca com logos
    "conteudo": 1,      # branco, logos no topo, faixa gradiente no rodapé
    "encerramento": 2,  # ondas + logos LAB365 e SENAI grandes no centro
    "logo": 3,          # ondas + logo LAB365 no centro
    "secao": 4,         # gradiente liso + logos no topo
    "faixa": 5,         # branco com faixa gradiente inferior
}


def cor(nome_ou_hex: str) -> str:
    """Resolve 'azul', 'texto_medio'... ou '#RRGGBB' para '#RRGGBB'."""
    if nome_ou_hex.startswith("#"):
        return nome_ou_hex
    return CORES[nome_ou_hex]


def rgb(nome_ou_hex: str) -> RGBColor:
    h = cor(nome_ou_hex).lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def emu(pol: float) -> Emu:
    return Emu(int(round(pol * EMU_POR_POL)))


# ── Acentos (paleta em sequência) ─────────────────────────────────────────────

@dataclass(frozen=True)
class Acento:
    """Trio de cores de um acento: forte (preenchimento), suave (fundo pastel) e tinta (texto/ícone sobre o pastel)."""
    nome: str
    forte: str
    suave: str
    tinta: str


ACENTOS = {n: Acento(n, a["forte"], a["suave"], a["tinta"]) for n, a in TOKENS["acentos"].items()}
SEQUENCIA = TOKENS["sequencia_acentos"]


def acento(chave: Union[int, str, None], padrao: int = 0) -> Acento:
    """Acento por índice (cicla laranja → magenta → roxo → azul) ou por nome."""
    if isinstance(chave, str) and chave in ACENTOS:
        return ACENTOS[chave]
    if isinstance(chave, int):
        return ACENTOS[SEQUENCIA[chave % len(SEQUENCIA)]]
    return ACENTOS[SEQUENCIA[padrao % len(SEQUENCIA)]]


def acento_item(item: dict, indice: int) -> Acento:
    """Acento de um item de lista: usa item['cor'] se vier no YAML, senão a sequência."""
    return acento(item.get("cor") if isinstance(item, dict) else None, indice)


# ── Medição de texto (Lexend) ─────────────────────────────────────────────────

_ARQ_FONTE = {
    ("Lexend", False): "Lexend-Regular.ttf",
    ("Lexend", True): "Lexend-Bold.ttf",
    ("Lexend Light", False): "Lexend-Light.ttf",
    ("Lexend Light", True): "Lexend-Bold.ttf",
    ("Lexend ExtraBold", False): "Lexend-ExtraBold.ttf",
    ("Lexend ExtraBold", True): "Lexend-ExtraBold.ttf",
}
_ESCALA = 20  # mede em 20 px por pt para precisão sub-ponto


@functools.lru_cache(maxsize=None)
def _fonte_pil(familia: str, negrito: bool) -> Optional[ImageFont.FreeTypeFont]:
    arq = _ARQ_FONTE.get((familia, negrito))
    if arq is None:
        return None
    return ImageFont.truetype(str(DIR_FONTES / arq), size=100 * _ESCALA)


@functools.lru_cache(maxsize=None)
def _altura_linha_em(familia: str) -> float:
    """Altura de uma linha em 'em' (ascendente + descendente + lineGap) como o PowerPoint calcula."""
    from fontTools.ttLib import TTFont  # import tardio: só na primeira medição

    arq = _ARQ_FONTE.get((familia, False), "Lexend-Regular.ttf")
    f = TTFont(str(DIR_FONTES / arq))
    hh = f["hhea"]
    return (hh.ascent - hh.descent + hh.lineGap) / f["head"].unitsPerEm


def largura_texto(texto: str, tam: float, negrito: bool = False, familia: str = "Lexend") -> float:
    """Largura em polegadas de uma linha de texto."""
    if familia == FONTES["codigo"]:
        return len(texto) * tam * 0.6 / 72  # Courier New é monoespaçada (600/1000 em)
    f = _fonte_pil(familia, negrito) or _fonte_pil("Lexend", negrito)
    return f.getlength(texto) / (100 * _ESCALA) * tam / 72


def altura_linha(tam: float, espac: float = 1.0, familia: str = "Lexend") -> float:
    """Altura em polegadas de uma linha de texto com espaçamento `espac` (1.0 = simples)."""
    em = 1.133 if familia == FONTES["codigo"] else _altura_linha_em(familia)
    return tam * em * espac / 72


def quebrar_linhas(texto: str, largura: float, tam: float, negrito: bool = False, familia: str = "Lexend") -> list[str]:
    """Quebra por palavra como o PowerPoint (aprox.). Respeita '\\n'."""
    linhas: list[str] = []
    for paragrafo in str(texto).split("\n"):
        palavras = paragrafo.split(" ")
        atual = ""
        for p in palavras:
            teste = p if not atual else f"{atual} {p}"
            if largura_texto(teste, tam, negrito, familia) <= largura or not atual:
                atual = teste
            else:
                linhas.append(atual)
                atual = p
        linhas.append(atual)
    return linhas


def altura_texto(texto: str, largura: float, tam: float, negrito: bool = False,
                 espac: float = 1.0, familia: str = "Lexend") -> float:
    """Altura em polegadas que o texto ocupa numa caixa de largura dada."""
    n = len(quebrar_linhas(texto, largura, tam, negrito, familia))
    return n * altura_linha(tam, espac, familia)


def cabe(texto: str, largura: float, altura: float, tam: float, negrito: bool = False,
         espac: float = 1.0, familia: str = "Lexend", max_linhas: Optional[int] = None) -> bool:
    linhas = quebrar_linhas(texto, largura, tam, negrito, familia)
    if max_linhas is not None and len(linhas) > max_linhas:
        return False
    # nenhuma palavra isolada pode ser mais larga que a caixa
    if any(largura_texto(l, tam, negrito, familia) > largura * 1.001 for l in linhas):
        return False
    return len(linhas) * altura_linha(tam, espac, familia) <= altura * 1.02


def tamanho_que_cabe(texto: str, largura: float, altura: float, tam: float, *, negrito: bool = False,
                     espac: float = 1.0, familia: str = "Lexend", minimo: Optional[float] = None,
                     max_linhas: Optional[int] = None) -> float:
    """Maior tamanho ≤ `tam` (passo de 0,5 pt) em que o texto cabe na caixa. Nunca abaixo de `minimo`."""
    minimo = minimo if minimo is not None else max(6.0, round(tam * 0.6 * 2) / 2)
    t = tam
    while t > minimo and not cabe(texto, largura, altura, t, negrito, espac, familia, max_linhas):
        t -= 0.5
    return max(t, minimo)


# ── Texto ─────────────────────────────────────────────────────────────────────

_ALINH = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
_ANCORA = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


@dataclass
class Trecho:
    """Um pedaço de texto com estilo próprio dentro de um parágrafo (ex.: destaque em negrito + texto normal)."""
    texto: str
    negrito: Optional[bool] = None
    cor: Optional[str] = None
    tam: Optional[float] = None
    familia: Optional[str] = None
    italico: bool = False


Paragrafo = Union[str, Sequence[Trecho]]


def _familia(nome: Optional[str]) -> str:
    if not nome:
        return FONTES["corpo"]
    return FONTES.get(nome, nome)


def _config_caixa(tf, ancora: str, ins: float = 0.0) -> None:
    tf.word_wrap = True
    tf.auto_size = None
    bp = tf._txBody.find(qn("a:bodyPr"))
    for lado in ("lIns", "tIns", "rIns", "bIns"):
        bp.set(lado, str(int(ins * EMU_POR_POL)))
    tf.vertical_anchor = _ANCORA[ancora]


def _escrever(tf, paragrafos: Sequence[Paragrafo], *, tam: float, cor_: str, negrito: bool,
              familia: str, alinh: str, espac: float, maiusc: bool, italico: bool = False,
              espaco_entre: float = 0.0) -> None:
    """Escreve parágrafos num text_frame vazio."""
    tf.clear()
    for i, par in enumerate(paragrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALINH[alinh]
        p.line_spacing = espac
        if i > 0 and espaco_entre:
            p.space_before = Pt(espaco_entre)
        trechos = [Trecho(par)] if isinstance(par, str) else list(par)
        for tr in trechos:
            r = p.add_run()
            txt = tr.texto.upper() if maiusc else tr.texto
            r.text = txt
            r._r.get_or_add_rPr().set("lang", "pt-BR")
            f = r.font
            f.name = _familia(tr.familia) if tr.familia else familia
            f.size = Pt(tr.tam if tr.tam else tam)
            f.bold = negrito if tr.negrito is None else tr.negrito
            f.italic = tr.italico or italico
            f.color.rgb = rgb(tr.cor or cor_)


def _texto_plano(paragrafos: Sequence[Paragrafo]) -> list[str]:
    return [p if isinstance(p, str) else "".join(t.texto for t in p) for p in paragrafos]


def texto(slide, x: float, y: float, w: float, h: float, conteudo: Union[Paragrafo, Sequence[Paragrafo]], *,
          estilo: Optional[str] = None, tam: Optional[float] = None, cor: Optional[str] = None,
          negrito: Optional[bool] = None, fonte: Optional[str] = None, alinh: str = "l", ancora: str = "t",
          espac: Optional[float] = None, maiusc: Optional[bool] = None, italico: bool = False,
          ajustar: bool = True, minimo: Optional[float] = None, max_linhas: Optional[int] = None,
          espaco_entre: float = 0.0, forma=None):
    """
    Caixa de texto no padrão do design system.

    `conteudo`: string (\\n separa parágrafos), lista de strings (um parágrafo cada)
    ou lista de listas de `Trecho` (texto rico).
    `estilo`: nome em tokens.estilos (eyebrow, titulo, rotulo, titulo_card, corpo...);
    parâmetros explícitos sobrescrevem o estilo.
    Com `ajustar=True` a fonte encolhe até caber na caixa (nunca estoura).
    `forma`: escreve dentro de uma forma existente em vez de criar caixa nova.
    """
    e = dict(ESTILOS.get(estilo, {})) if estilo else {}
    tam = tam if tam is not None else e.get("tam", 10)
    cor_ = cor if cor is not None else e.get("cor", "texto_medio")
    negrito = negrito if negrito is not None else e.get("negrito", False)
    familia = _familia(fonte if fonte is not None else e.get("fonte"))
    espac = espac if espac is not None else e.get("espac", 1.0)
    maiusc = maiusc if maiusc is not None else e.get("maiusc", False)

    if isinstance(conteudo, str):
        paragrafos: list[Paragrafo] = conteudo.split("\n")
    elif conteudo and all(isinstance(t, Trecho) for t in conteudo):
        paragrafos = [list(conteudo)]
    else:
        paragrafos = list(conteudo)

    if ajustar:
        plano = "\n".join(_texto_plano(paragrafos))
        if maiusc:
            plano = plano.upper()
        tam_ok = tamanho_que_cabe(plano, w, h, tam, negrito=negrito, espac=espac, familia=familia,
                                  minimo=minimo, max_linhas=max_linhas)
        if tam_ok < tam:
            fator = tam_ok / tam
            tam = tam_ok
            paragrafos = [p if isinstance(p, str) else
                          [Trecho(t.texto, t.negrito, t.cor, (t.tam * fator if t.tam else None), t.familia, t.italico) for t in p]
                          for p in paragrafos]

    if forma is None:
        forma = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = forma.text_frame
    _config_caixa(tf, ancora)
    _escrever(tf, paragrafos, tam=tam, cor_=cor_, negrito=negrito, familia=familia, alinh=alinh,
              espac=espac, maiusc=maiusc, italico=italico, espaco_entre=espaco_entre)
    return forma


# ── Formas ────────────────────────────────────────────────────────────────────

_FORMAS = {
    "ret": MSO_SHAPE.RECTANGLE,
    "arred": MSO_SHAPE.ROUNDED_RECTANGLE,
    "elipse": MSO_SHAPE.OVAL,
    "chevron": MSO_SHAPE.CHEVRON,
    "seta": MSO_SHAPE.PENTAGON,        # homePlate: primeira etapa de um fluxo
    "arco": MSO_SHAPE.BLOCK_ARC,
    "triangulo": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "seta_direita": MSO_SHAPE.RIGHT_ARROW,
}


def forma(slide, tipo: str, x: float, y: float, w: float, h: float, *, preench: Optional[str] = None,
          borda: Optional[str] = None, borda_larg: float = 1.0, tracejado: bool = False,
          adj: Optional[Sequence[float]] = None, rot: float = 0.0):
    """Forma geométrica sem sombra. `preench`/`borda` aceitam nome de token ou #hex; None = sem."""
    s = slide.shapes.add_shape(_FORMAS[tipo], emu(x), emu(y), emu(w), emu(h))
    if preench:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(preench)
    else:
        s.fill.background()
    if borda:
        s.line.color.rgb = rgb(borda)
        s.line.width = Pt(borda_larg)
        if tracejado:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    estilo_tema = s._element.find(qn("p:style"))
    if estilo_tema is not None:  # sem p:style: nada de sombra/contorno herdado do tema
        s._element.remove(estilo_tema)
    if adj:
        for i, v in enumerate(adj):
            s.adjustments[i] = v
    if rot:
        s.rotation = rot
    # texto interno herdado do tema: zera para não aparecer "Arial 18"
    s.text_frame.text = ""
    return s


def card(slide, x: float, y: float, w: float, h: float, preench: str = "fundo_neutro", raio: float = RAIO, **kw):
    """Retângulo arredondado com raio absoluto (em polegadas), igual em qualquer tamanho."""
    adj = min(0.5, raio / max(min(w, h), 0.01))
    return forma(slide, "arred", x, y, w, h, preench=preench, adj=[adj], **kw)


def circulo(slide, x: float, y: float, d: float, preench: Optional[str] = None, **kw):
    return forma(slide, "elipse", x, y, d, d, preench=preench, **kw)


def linha(slide, x1: float, y1: float, x2: float, y2: float, cor_: str = "borda", larg: float = 1.0,
          tracejado: bool = False, seta_fim: bool = False):
    """Conector reto de (x1,y1) a (x2,y2)."""
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    c.line.color.rgb = rgb(cor_)
    c.line.width = Pt(larg)
    if tracejado:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if seta_fim:
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return c


def pill(slide, x: float, y: float, w: float, h: float, rotulo: str, *, preench: str, cor_texto: str,
         tam: Optional[float] = None):
    """Etiqueta arredondada (cápsula) com texto centralizado em caixa alta."""
    s = forma(slide, "arred", x, y, w, h, preench=preench, adj=[0.5])
    e = ESTILOS["pill"]
    texto(slide, x, y, w, h, rotulo, tam=tam or e["tam"], negrito=True, cor=cor_texto, maiusc=True,
          alinh="c", ancora="m", max_linhas=1, minimo=5.5, forma=s)
    return s


def largura_pill(rotulo: str, tam: Optional[float] = None, folga: float = 0.28) -> float:
    """Largura ideal de uma pill para o rótulo (texto + folga lateral)."""
    t = tam or ESTILOS["pill"]["tam"]
    return largura_texto(rotulo.upper(), t, True) + folga


# ── Ícones (Lucide) ───────────────────────────────────────────────────────────

ICONE_PADRAO = "sparkles"


def _kebab(nome: str) -> str:
    n = str(nome).strip()
    n = re.sub(r"\.(png|svg)$", "", n, flags=re.I)
    n = re.sub(r"_[0-9A-Fa-f]{6}$", "", n)               # 'Target_0E1D8E' → 'Target'
    n = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", "-", n)          # CircleCheck → Circle-Check
    n = re.sub(r"[\s_]+", "-", n)
    return n.lower()


@functools.lru_cache(maxsize=1)
def catalogo_icones() -> dict[str, list[str]]:
    """Nome kebab-case → tags (em inglês) de todos os ícones disponíveis."""
    arq = DIR_ICONES / "catalogo.json"
    if arq.is_file():
        with open(arq, encoding="utf-8") as f:
            return json.load(f)
    return {p.stem: [] for p in DIR_ICONES.glob("*.png")}


def resolver_icone(nome: Optional[str]) -> Optional[str]:
    """Nome canônico (kebab-case) do ícone ou None se não existir. Aceita 'CircleCheck', 'circle-check', 'Target_0E1D8E.png'."""
    if not nome:
        return None
    k = _kebab(nome)
    return k if k in catalogo_icones() else None


def sugerir_icones(termo: str, n: int = 6) -> list[str]:
    """Ícones parecidos com `termo` (por nome e por tag)."""
    cat = catalogo_icones()
    k = _kebab(termo)
    achados = difflib.get_close_matches(k, cat.keys(), n=n, cutoff=0.5)
    palavras = set(re.split(r"[-\s]+", k))
    for nome, tags in cat.items():
        if len(achados) >= n:
            break
        if nome not in achados and (palavras & set(tags) or any(p in nome.split("-") for p in palavras)):
            achados.append(nome)
    return achados[:n]


@functools.lru_cache(maxsize=1024)
def _png_icone(nome: str, cor_hex: str) -> bytes:
    """PNG do ícone recolorido (a máscara branca vira a cor pedida mantendo o alpha)."""
    mascara = Image.open(DIR_ICONES / f"{nome}.png").convert("RGBA")
    alpha = mascara.getchannel("A")
    h = cor_hex.lstrip("#")
    solido = Image.new("RGBA", mascara.size, (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16), 255))
    solido.putalpha(alpha)
    buf = io.BytesIO()
    solido.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def icone(slide, nome: Optional[str], x: float, y: float, tam: float, cor_: str = "branco"):
    """Ícone Lucide recolorido, quadrado de lado `tam`. Ícone inexistente cai no padrão (sparkles)."""
    canon = resolver_icone(nome) or ICONE_PADRAO
    png = _png_icone(canon, cor(cor_))
    pic = slide.shapes.add_picture(io.BytesIO(png), emu(x), emu(y), emu(tam), emu(tam))
    pic.name = f"icone:{canon}"
    return pic


def icone_circulo(slide, nome: Optional[str], x: float, y: float, d: float, *, fundo: str,
                  cor_icone: str = "branco", proporcao: float = 0.54, borda: Optional[str] = None):
    """Ícone centralizado num círculo de diâmetro `d` (padrão das bolinhas do deck oficial)."""
    circulo(slide, x, y, d, preench=fundo, borda=borda)
    t = d * proporcao
    return icone(slide, nome, x + (d - t) / 2, y + (d - t) / 2, t, cor_icone)


# ── Deck e slides ─────────────────────────────────────────────────────────────

class Deck:
    """Apresentação aberta a partir do template oficial LAB365 (layouts, logos, fundos e fontes embutidas)."""

    def __init__(self, template: Union[str, Path] = TEMPLATE) -> None:
        self.prs = Presentation(str(template))

    def novo_slide(self, layout: str = "conteudo"):
        """Adiciona slide com o layout pedido; remove placeholders exceto o título."""
        s = self.prs.slides.add_slide(self.prs.slide_layouts[LAYOUTS[layout]])
        for ph in list(s.placeholders):
            if ph.placeholder_format.type not in (1, 3):  # TITLE, CENTER_TITLE
                ph._element.getparent().remove(ph._element)
        return s

    def salvar(self, caminho: Union[str, Path]) -> None:
        Path(caminho).parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(caminho))


def _placeholder_titulo(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.type in (1, 3):
            return ph
    return None


def titulo(slide, conteudo: str, x: float, y: float, w: float, h: float, *, estilo: str = "titulo",
           alinh: str = "l", ancora: str = "m", max_linhas: int = 1, minimo: Optional[float] = None, **kw):
    """Escreve o título no placeholder de título do layout (mantém a estrutura/acessibilidade do PPTX)."""
    ph = _placeholder_titulo(slide)
    if ph is None:
        return texto(slide, x, y, w, h, conteudo, estilo=estilo, alinh=alinh, ancora=ancora,
                     max_linhas=max_linhas, minimo=minimo, **kw)
    ph.left, ph.top, ph.width, ph.height = emu(x), emu(y), emu(w), emu(h)
    return texto(slide, x, y, w, h, conteudo, estilo=estilo, alinh=alinh, ancora=ancora,
                 max_linhas=max_linhas, minimo=minimo, forma=ph, **kw)


def cabecalho(slide, eyebrow: str, titulo_: str) -> None:
    """Eyebrow magenta + título azul em caixa alta, no mesmo lugar em todo slide de conteúdo."""
    if eyebrow:
        texto(slide, MX, GRID["topo_eyebrow"], 6.5, 0.2, eyebrow, estilo="eyebrow", max_linhas=1, minimo=6)
    titulo(slide, titulo_, MX, GRID["topo_titulo"], GRID["largura_titulo"], GRID["altura_titulo"],
           estilo="titulo", max_linhas=1, minimo=15)


def separar_destaque(txt: str) -> tuple[str, str]:
    """'Regra do jogo: texto' → ('Regra do jogo: ', 'texto'). Sem ':' → ('', txt)."""
    m = re.match(r"^(.{2,40}?:)\s*(.+)$", txt.strip(), flags=re.S)
    if m:
        return m.group(1) + " ", m.group(2)
    return "", txt.strip()


def nota(slide, dados: Union[str, dict, None], *, y: float = TOPO_NOTA, h: float = ALT_NOTA,
         x: float = MX, w: float = LU, cor_padrao: str = "roxo") -> None:
    """
    Barra de destaque no rodapé (pastel + bolinha com ícone + "Destaque: texto").

    `dados`: "Destaque: texto" ou {texto, destaque?, icone?, cor?}.
    """
    if not dados:
        return
    if isinstance(dados, str):
        dados = {"texto": dados}
    a = acento(dados.get("cor") or cor_padrao)
    destaque = dados.get("destaque")
    corpo = dados.get("texto", "")
    if destaque is None:
        destaque, corpo = separar_destaque(corpo)
    elif destaque and not destaque.endswith(" "):
        destaque = destaque.rstrip(":") + ": "
    card(slide, x, y, w, h, preench=a.suave)
    d = min(0.5, h - 0.2)
    icone_circulo(slide, dados.get("icone") or "lightbulb", x + 0.2, y + (h - d) / 2, d, fundo=a.forte,
                  proporcao=0.54)
    tx = x + 0.2 + d + 0.25
    texto(slide, tx, y, x + w - 0.25 - tx, h,
          [[Trecho(destaque, negrito=True, cor=a.tinta), Trecho(corpo)]] if destaque else corpo,
          estilo="nota", ancora="m", max_linhas=2, minimo=8.5)


def area_conteudo(tem_nota: bool) -> tuple[float, float]:
    """(topo, base) livres para o corpo do slide, descontando a barra de destaque se houver."""
    return TOPO, (TOPO_NOTA - ESPACO if tem_nota else BASE)


def notas_apresentador(slide, txt: Optional[str]) -> None:
    """Texto para o apresentador (painel de anotações do PowerPoint)."""
    if txt:
        slide.notes_slide.notes_text_frame.text = str(txt)


def distribuir(n: int, x: float, w: float, gap: float) -> list[tuple[float, float]]:
    """Divide a largura `w` a partir de `x` em `n` colunas com espaço `gap`. Retorna [(x_i, largura)]."""
    larg = (w - gap * (n - 1)) / n
    return [(x + i * (larg + gap), larg) for i in range(n)]


def grade(n: int, colunas: int, x: float, y: float, w: float, h: float, gap_x: float, gap_y: float):
    """Posições (x, y, w, h) de `n` células numa grade com `colunas` colunas preenchendo a área."""
    linhas = -(-n // colunas)
    cw = (w - gap_x * (colunas - 1)) / colunas
    ch = (h - gap_y * (linhas - 1)) / linhas
    return [(x + (i % colunas) * (cw + gap_x), y + (i // colunas) * (ch + gap_y), cw, ch) for i in range(n)]
